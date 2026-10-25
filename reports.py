from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import storage as db

WD = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август",
          "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
MON_SHORT = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]
KIND_NAMES = {"tip": "Чаевые", "cash": "Наличка", "car": "Машина", "fuel": "Заправка", "other": "Прочее"}
KEYS = ("yandex", "nal", "tips", "car", "fuel", "other", "exp", "minutes", "km", "orders", "cash", "litres")


def fmt(x):
    x = round(x or 0, 2)
    if x == int(x):
        return str(int(x))
    return f"{x:.2f}".rstrip("0")


def dur(minutes):
    minutes = int(minutes)
    return f"{minutes // 60:02d}ч{minutes % 60:02d}мин"


def month_label(ym):
    y, m = ym.split("-")
    return f"{MONTHS[int(m) - 1]} {y}"


def calc(s, entries):
    tips = sum(e["amount"] for e in entries if e["kind"] == "tip")
    nal = sum(e["amount"] for e in entries if e["kind"] == "cash")
    by = {k: sum(e["amount"] for e in entries if e["kind"] == k) for k in ("car", "fuel", "other")}
    exp = sum(by.values())
    yandex = s["yandex"] or 0
    end_ts = s["end_ts"] or s["start_ts"]
    km = (s["end_odo"] - s["start_odo"]) if s["end_odo"] is not None else 0
    return {
        "yandex": yandex, "nal": nal, "tips": tips, "car": by["car"], "fuel": by["fuel"], "other": by["other"],
        "exp": exp, "minutes": max(0, (end_ts - s["start_ts"]) // 60), "km": km,
        "orders": s["orders"] or 0, "cash": yandex + nal + tips - exp,
        "litres": sum(e["litres"] or 0 for e in entries if e["kind"] == "fuel"),
    }


def total(calcs):
    return {k: sum(c[k] for c in calcs) for k in KEYS}


def _exp_line(label, kind, entries, c):
    text = f"- {label}: {fmt(c[kind])}"
    notes = [f"{e['note']} {fmt(e['amount'])}" for e in entries if e["kind"] == kind and e["note"]]
    if notes:
        text += " (" + ", ".join(notes) + ")"
    if kind == "fuel" and c["litres"]:
        text += f" · {fmt(c['litres'])} л"
    return text


def _exp_block(entries, c):
    return [_exp_line("Машина", "car", entries, c),
            _exp_line("Заправка", "fuel", entries, c),
            _exp_line("Прочее", "other", entries, c)]


def day_report(s, entries, n, plan, month_cash):
    c = calc(s, entries)
    d = db.dt(s["start_ts"])
    head = f"🗓️ Итоги смены {n}/{plan}" if plan else f"🗓️ Итоги смены {n}"
    head += f": {WD[d.weekday()]} {d:%d.%m.%y}"
    lines = [head, "",
             f"🎰 Касса месяца: {fmt(month_cash)}",
             f"🧾 Касса дня: {fmt(c['cash'])}", "",
             f"🟢 На линии: {dur(c['minutes'])}",
             f"🚕 Пробег: {fmt(c['km'])}",
             f"🙋🏻 Заказов: {c['orders']}",
             f"🪙 Наличные: {fmt(c['nal'])}",
             f"🫴 Яндекс: {fmt(c['yandex'])}",
             f"🤑 Чаевые: {fmt(c['tips'])}", "",
             "📉 Расходы:"] + _exp_block(entries, c)
    return "\n".join(lines)


def status_text(s, entries):
    c = calc(s, entries)
    d = db.dt(s["start_ts"])
    lines = [f"🟢 Смена открыта: {d:%d.%m.%y %H:%M}",
             f"🚕 Пробег на старте: {fmt(s['start_odo'])}",
             f"🪙 Наличные: {fmt(c['nal'])}",
             f"🤑 Чаевые: {fmt(c['tips'])}",
             f"📉 Расходы: {fmt(c['exp'])}"]
    return "\n".join(lines)


def entries_text(entries):
    if not entries:
        return "Записей пока нет."
    out = []
    for i, e in enumerate(entries, 1):
        t = db.dt(e["ts"]).strftime("%H:%M")
        extra = f" ({e['note']})" if e["note"] else ""
        lit = f" · {fmt(e['litres'])} л" if e["litres"] else ""
        out.append(f"{i}. {KIND_NAMES[e['kind']]}: {fmt(e['amount'])}{extra}{lit} · {t}")
    return "\n".join(out)


def period_report(kind, title, calcs, by_month=None):
    n = len(calcs)
    if n == 0:
        return f"За {title} закрытых смен пока нет."
    t = total(calcs)
    word = "месяца" if kind == "month" else "года"
    lines = [f"🗓️ Итоги {word}: {title}", "",
             f"🎰 Касса {word}: {fmt(t['cash'])}",
             f"📊 Смен: {n} (в среднем {fmt(round(t['cash'] / n))} за смену)", "",
             f"🟢 На линии: {dur(t['minutes'])}",
             f"🚕 Пробег: {fmt(t['km'])}",
             f"🙋🏻 Заказов: {t['orders']}",
             f"🪙 Наличные: {fmt(t['nal'])}",
             f"🫴 Яндекс: {fmt(t['yandex'])}",
             f"🤑 Чаевые: {fmt(t['tips'])}", "",
             f"📉 Расходы: {fmt(t['exp'])}",
             f"- Машина: {fmt(t['car'])}",
             f"- Заправка: {fmt(t['fuel'])}",
             f"- Прочее: {fmt(t['other'])}"]
    extra = []
    if t["km"] > 0:
        extra.append(f"💡 Заработок: {fmt(round(t['cash'] / t['km'], 1))} ₽ за км")
        if t["litres"] > 0:
            extra.append(f"⛽ Расход: {fmt(round(t['litres'] / t['km'] * 100, 1))} л на 100 км")
    if extra:
        lines += [""] + extra
    if by_month:
        lines += ["", "📅 По месяцам (касса):"]
        cells = [f"{MON_SHORT[int(ym[5:]) - 1]} {fmt(sum(c['cash'] for c in cs))}"
                 for ym, cs in sorted(by_month.items())]
        for i in range(0, len(cells), 3):
            lines.append(" · ".join(cells[i:i + 3]))
    return "\n".join(lines)


def build_xlsx(title, items, ents, by_month=None):
    """items: [(shift, calc)], ents: {shift_id: [entries]}, by_month: {ym: [calc]}."""
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Смены"
    head = ["№", "Дата", "Начало", "Конец", "На линии (ч)", "Пробег", "Заказов", "Наличные",
            "Яндекс", "Чаевые", "Машина", "Заправка", "Прочее", "Расходы", "Касса дня"]
    ws.append(head)
    for i, (s, c) in enumerate(items, 1):
        st, en = db.dt(s["start_ts"]), db.dt(s["end_ts"])
        ws.append([i, st.strftime("%d.%m.%Y"), st.strftime("%H:%M"), en.strftime("%H:%M"),
                   round(c["minutes"] / 60, 2), c["km"], c["orders"], c["nal"], c["yandex"], c["tips"],
                   c["car"], c["fuel"], c["other"], c["exp"], c["cash"]])
    if items:
        last = len(items) + 1
        ws.append(["Итого", "", "", ""] + [f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{last})"
                                          for col in range(5, 16)])
        for cell in ws[last + 1]:
            cell.font = bold
    _style(ws, head)

    ws2 = wb.create_sheet("Записи")
    head2 = ["Дата", "Время", "Тип", "Сумма", "Пояснение", "Литры"]
    ws2.append(head2)
    for s, _ in items:
        for e in ents.get(s["id"], []):
            d = db.dt(e["ts"])
            ws2.append([d.strftime("%d.%m.%Y"), d.strftime("%H:%M"), KIND_NAMES[e["kind"]],
                        e["amount"], e["note"] or "", e["litres"] or ""])
    _style(ws2, head2)

    if by_month:
        ws3 = wb.create_sheet("По месяцам")
        head3 = ["Месяц", "Смен", "Пробег", "Наличные", "Яндекс", "Чаевые", "Расходы", "Касса"]
        ws3.append(head3)
        for ym, cs in sorted(by_month.items()):
            t = total(cs)
            ws3.append([month_label(ym), len(cs), t["km"], t["nal"], t["yandex"], t["tips"], t["exp"], t["cash"]])
        last = len(by_month) + 1
        ws3.append(["Итого"] + [f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{last})"
                                for col in range(2, 9)])
        for cell in ws3[last + 1]:
            cell.font = bold
        _style(ws3, head3)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style(ws, head):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = "A2"
